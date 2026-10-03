"""Word reports end to end: options, templates, generation with evidence."""
import io
import json
import uuid
from pathlib import Path

import pytest
from docx import Document

from scanr.reporting import docx_renderer

FIXTURES = Path(__file__).parents[1] / "fixtures" / "imports"


@pytest.fixture(autouse=True)
def storage(tmp_path, monkeypatch):
    from scanr.config import get_settings

    settings = get_settings()
    monkeypatch.setattr(settings, "reports_dir", tmp_path / "reports")
    monkeypatch.setattr(settings, "evidence_dir", tmp_path / "evidence")
    queued = []
    from scanr.tasks import report_tasks

    monkeypatch.setattr(report_tasks.generate_report_task, "delay", lambda report_id: queued.append(report_id))
    return queued


async def _scan_with_findings(client, headers):
    r = await client.post("/api/v1/scans/import", headers=headers,
                          json={"name": f"docx {uuid.uuid4().hex[:6]}", "report": (FIXTURES / "sample.nessus").read_text()})
    return r.json()["scan_id"]


@pytest.mark.asyncio
async def test_create_docx_report_and_generate_it(client, auth_headers, db, storage):
    scan_id = await _scan_with_findings(client, auth_headers)
    findings = (await client.get(f"/api/v1/findings?scan_id={scan_id}&sort=priority", headers=auth_headers)).json()
    png = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde"
           b"\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0\x00\x00\x03\x01\x01\x00\xc9\xfe\x92\xef\x00\x00\x00\x00IEND\xaeB`\x82")
    await client.post(f"/api/v1/findings/{findings[0]['id']}/attachments", headers=auth_headers,
                      files={"file": ("shot.png", png, "image/png")}, data={"caption": "Exploited with Metasploit"})

    r = await client.post("/api/v1/reports", headers=auth_headers, json={
        "scan_id": scan_id, "format": "docx", "title": "Internal test", "client": "ACME", "author": "Tester"})
    assert r.status_code == 201, r.text
    report_id = r.json()["id"]
    assert storage == [report_id]

    from scanr.models import Report
    from scanr.reporting.report_engine import ReportEngine

    report = await db.get(Report, report_id)
    assert json.loads(report.options) == {"title": "Internal test", "client": "ACME", "author": "Tester", "include_info": False}
    out = await ReportEngine(db).generate(report)
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Internal test" in text and "ACME" in text and "Prepared by: Tester" in text
    assert "F-01  MS17-010" in text and "Exploited with Metasploit" in text
    assert "Nessus Scan Information" not in text  # informational hidden by default
    assert len(doc.inline_shapes) == 1
    overview = doc.tables[1]
    assert overview.rows[1].cells[2].text == "CRITICAL"


@pytest.mark.asyncio
async def test_template_upload_permissions_and_use(client, auth_headers, db):
    default = await client.get("/api/v1/report-templates/default/download", headers=auth_headers)
    assert default.status_code == 200 and default.content[:2] == b"PK"
    branded = Document(io.BytesIO(default.content))
    branded.paragraphs[0].text = "ACME Security Ltd — {{ report.client }}"
    buf = io.BytesIO()
    branded.save(buf)

    name = f"House style {uuid.uuid4().hex[:4]}"
    up = await client.post("/api/v1/report-templates", headers=auth_headers, data={"name": name},
                           files={"file": ("house.docx", buf.getvalue(), "application/octet-stream")})
    assert up.status_code == 201, up.text
    tid = up.json()["id"]
    dup = await client.post("/api/v1/report-templates", headers=auth_headers, data={"name": name.upper()},
                            files={"file": ("house.docx", buf.getvalue(), "application/octet-stream")})
    assert dup.status_code == 409
    bad = await client.post("/api/v1/report-templates", headers=auth_headers, data={"name": "broken"},
                            files={"file": ("x.docx", b"not a docx", "application/octet-stream")})
    assert bad.status_code == 400

    email = f"docx-analyst-{uuid.uuid4().hex[:4]}@scanr.local"
    await client.post("/api/v1/users", json={"email": email, "password": "long-enough-pw"}, headers=auth_headers)
    login = await client.post("/api/v1/auth/login", json={"email": email, "password": "long-enough-pw"})
    analyst = {"Authorization": f"Bearer {login.json()['access_token']}"}
    assert (await client.get("/api/v1/report-templates", headers=analyst)).status_code == 200
    assert (await client.post("/api/v1/report-templates", headers=analyst, data={"name": "x"},
                              files={"file": ("x.docx", buf.getvalue(), "application/octet-stream")})).status_code == 403

    scan_id = await _scan_with_findings(client, auth_headers)
    r = await client.post("/api/v1/reports", headers=auth_headers,
                          json={"scan_id": scan_id, "format": "docx", "template_id": tid, "client": "Globex"})
    from scanr.models import Report
    from scanr.reporting.report_engine import ReportEngine

    out = await ReportEngine(db).generate(await db.get(Report, r.json()["id"]))
    assert Document(str(out)).paragraphs[0].text == "ACME Security Ltd — Globex"

    missing = await client.post("/api/v1/reports", headers=auth_headers,
                                json={"scan_id": scan_id, "format": "docx", "template_id": "nope"})
    assert missing.status_code == 404
    assert (await client.delete(f"/api/v1/report-templates/{tid}", headers=auth_headers)).status_code == 204
    assert not (docx_renderer.template_dir() / f"{tid}.docx").exists()


@pytest.mark.asyncio
async def test_placeholder_help(client, auth_headers):
    r = await client.get("/api/v1/report-templates/placeholders", headers=auth_headers)
    assert "findings (list)" in r.json()["placeholders"]
